import io
import math
import os
from datetime import datetime, timedelta
from pathlib import Path

from flask import Blueprint, jsonify, request, send_file, url_for
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors
from werkzeug.utils import secure_filename

from models import (
    ColaboradorEvento,
    Convidado,
    DespesaEvento,
    Evento,
    FotoEvento,
    Fornecedor,
    ItemCompartilhado,
    ItemEvento,
    LogEvento,
    Notificacao,
    ProdutoFornecedor,
    TarefaEvento,
    Usuario,
    db,
)
from utils.auth import acesso_evento, login_required, usuario_atual

evento_detalhes_controller = Blueprint("evento_detalhes_controller", __name__)


def _erro(evento):
    if evento is None:
        return jsonify({"erro": "Evento não encontrado."}), 404
    if evento is False:
        return jsonify({"erro": "Você não tem acesso a este evento."}), 403
    return None


def _log(evento_id, acao, detalhe=""):
    usuario = usuario_atual()
    db.session.add(LogEvento(evento_id=evento_id, usuario_id=usuario.id if usuario else None, acao=acao, detalhe=detalhe))


def _recalcular_progresso(evento):
    if not evento.tarefas_rel:
        return
    feitas = sum(1 for t in evento.tarefas_rel if t.feita)
    evento.progresso = round((feitas / len(evento.tarefas_rel)) * 100)


# ---------------- Convidados / RSVP ----------------

@evento_detalhes_controller.get("/eventos/<int:evento_id>/convidados")
@login_required
def listar_convidados(evento_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    return jsonify([c.to_dict() for c in evento.convidados_rel])


@evento_detalhes_controller.post("/eventos/<int:evento_id>/convidados")
@login_required
def adicionar_convidado(evento_id):
    usuario, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    nome = (dados.get("nome") or "").strip()
    contato = (dados.get("contato") or "").strip()
    if not nome:
        return jsonify({"erro": "Nome do convidado é obrigatório."}), 400

    email = (dados.get("email") or "").strip().lower()
    telefone = (dados.get("telefone") or "").strip()
    if contato:
        if "@" in contato and not email:
            email = contato.lower()
        elif not telefone:
            telefone = contato

    convidado = Convidado(
        evento_id=evento.id,
        nome=nome,
        email=email,
        telefone=telefone,
        vinculo=(dados.get("vinculo") or "").strip(),
        status=(dados.get("status") or "pendente").strip(),
        restricoes_alimentares=(dados.get("restricoes") or dados.get("restricoes_alimentares") or "").strip(),
    )
    db.session.add(convidado)
    db.session.flush()
    _log(evento.id, "convidado_adicionado", convidado.nome)
    db.session.commit()
    resposta = convidado.to_dict()
    resposta["link_rsvp"] = f"/rsvp.html?token={convidado.token_rsvp}"
    return jsonify(resposta), 201


@evento_detalhes_controller.put("/eventos/<int:evento_id>/convidados/<int:convidado_id>")
@login_required
def atualizar_convidado(evento_id, convidado_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    convidado = Convidado.query.filter_by(id=convidado_id, evento_id=evento.id).first()
    if not convidado:
        return jsonify({"erro": "Convidado não encontrado."}), 404
    dados = request.get_json(silent=True) or {}
    for campo in ("nome", "email", "telefone", "vinculo"):
        if campo in dados:
            setattr(convidado, campo, (dados.get(campo) or "").strip())
    if "status" in dados and dados["status"] in {"pendente", "confirmado", "recusado"}:
        convidado.status = dados["status"]
    if "restricoes" in dados or "restricoes_alimentares" in dados:
        convidado.restricoes_alimentares = (dados.get("restricoes") or dados.get("restricoes_alimentares") or "").strip()
    _log(evento.id, "convidado_atualizado", convidado.nome)
    db.session.commit()
    return jsonify(convidado.to_dict())


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/convidados/<int:convidado_id>")
@login_required
def remover_convidado(evento_id, convidado_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    convidado = Convidado.query.filter_by(id=convidado_id, evento_id=evento.id).first()
    if not convidado:
        return jsonify({"erro": "Convidado não encontrado."}), 404
    nome = convidado.nome
    db.session.delete(convidado)
    _log(evento.id, "convidado_removido", nome)
    db.session.commit()
    return "", 204


@evento_detalhes_controller.get("/rsvp/<token>")
def ver_rsvp(token):
    convidado = Convidado.query.filter_by(token_rsvp=token).first()
    if not convidado:
        return jsonify({"erro": "Convite inválido ou expirado."}), 404
    evento = convidado.evento
    confirmados_publicos = []
    if evento.privacidade_convidados == "todos":
        confirmados_publicos = [c.nome for c in evento.convidados_rel if c.status == "confirmado"]
    return jsonify({
        "convidado": convidado.to_dict(incluir_token=False),
        "evento": {
            "nome": evento.nome_evento,
            "tipo": evento.tipo_evento,
            "data": evento.data_evento.isoformat(),
            "hora": evento.hora_evento or "",
            "local": evento.endereco_evento or "",
            "organizador": evento.owner.nome if evento.owner else "Momentus",
            "confirmados_publicos": confirmados_publicos,
            "itens_compartilhados": [i.to_dict() for i in evento.itens_compartilhados_rel],
        },
    })


@evento_detalhes_controller.post("/rsvp/<token>")
def responder_rsvp(token):
    convidado = Convidado.query.filter_by(token_rsvp=token).first()
    if not convidado:
        return jsonify({"erro": "Convite inválido ou expirado."}), 404
    dados = request.get_json(silent=True) or {}
    status = dados.get("status")
    if status not in {"confirmado", "recusado", "pendente"}:
        return jsonify({"erro": "Status de RSVP inválido."}), 400
    convidado.status = status
    if "restricoes" in dados:
        convidado.restricoes_alimentares = (dados.get("restricoes") or "").strip()
    db.session.add(Notificacao(
        usuario_id=convidado.evento.owner_id,
        tipo="rsvp",
        titulo="Resposta de convite",
        mensagem=f"{convidado.nome} atualizou a presença para {status}.",
        evento_id=convidado.evento_id,
        referencia_id=convidado.id,
    ))
    db.session.commit()
    return jsonify({"convidado": convidado.to_dict(incluir_token=False)})


@evento_detalhes_controller.post("/rsvp/<token>/itens/<int:item_id>/assumir")
def assumir_item_rsvp(token, item_id):
    convidado = Convidado.query.filter_by(token_rsvp=token).first()
    if not convidado:
        return jsonify({"erro": "Convite inválido ou expirado."}), 404
    item = ItemCompartilhado.query.filter_by(id=item_id, evento_id=convidado.evento_id).first()
    if not item:
        return jsonify({"erro": "Item não encontrado."}), 404
    if item.responsavel_nome and not (item.responsavel_tipo == "convidado" and item.responsavel_id == convidado.id):
        return jsonify({"erro": f"{item.responsavel_nome} já ficou responsável por este item."}), 409
    item.responsavel_tipo = "convidado"
    item.responsavel_id = convidado.id
    item.responsavel_nome = convidado.nome
    db.session.add(Notificacao(
        usuario_id=convidado.evento.owner_id, tipo="item-assumido", titulo="Item da lista assumido",
        mensagem=f"{convidado.nome} ficou responsável por {item.nome}.", evento_id=convidado.evento_id, referencia_id=item.id,
    ))
    db.session.commit()
    return jsonify(item.to_dict())


@evento_detalhes_controller.post("/rsvp/<token>/itens/<int:item_id>/liberar")
def liberar_item_rsvp(token, item_id):
    convidado = Convidado.query.filter_by(token_rsvp=token).first()
    if not convidado:
        return jsonify({"erro": "Convite inválido ou expirado."}), 404
    item = ItemCompartilhado.query.filter_by(id=item_id, evento_id=convidado.evento_id).first()
    if not item:
        return jsonify({"erro": "Item não encontrado."}), 404
    if not (item.responsavel_tipo == "convidado" and item.responsavel_id == convidado.id):
        return jsonify({"erro": "Você só pode liberar um item que assumiu."}), 403
    item.responsavel_tipo = None
    item.responsavel_id = None
    item.responsavel_nome = ""
    db.session.commit()
    return jsonify(item.to_dict())


# ---------------- Colaboradores ----------------

@evento_detalhes_controller.get("/colaboracoes/<token>")
def ver_colaboracao(token):
    """Exibe somente os dados necessários para a tela pública do convite."""
    colab = ColaboradorEvento.query.filter_by(token=token).first()
    if not colab:
        return jsonify({"erro": "Convite de colaboração inválido ou expirado."}), 404
    return jsonify({
        "colaborador": {
            "nome": colab.nome,
            "email": colab.email,
            "status": colab.status,
            "token": colab.token,
        },
        "evento": {
            "id": colab.evento.id,
            "nome": colab.evento.nome_evento,
            "data": colab.evento.data_evento.isoformat() if colab.evento.data_evento else None,
            "hora": colab.evento.hora_evento or "",
            "organizador": colab.evento.owner.nome if colab.evento.owner else "Momentus",
        },
    })

@evento_detalhes_controller.post("/eventos/<int:evento_id>/colaboradores")
@login_required
def adicionar_colaborador(evento_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    if evento.owner_id != usuario.id:
        return jsonify({"erro": "Somente a pessoa dona do evento pode convidar colaboradores."}), 403

    dados = request.get_json(silent=True) or {}
    nome = (dados.get("nome") or "").strip()
    email = (dados.get("email") or dados.get("contato") or "").strip().lower()
    if not nome or "@" not in email:
        return jsonify({"erro": "Informe nome e e-mail do colaborador."}), 400
    if email == usuario.email.lower():
        return jsonify({"erro": "Você já é a pessoa dona deste evento."}), 400
    if ColaboradorEvento.query.filter_by(evento_id=evento.id, email=email).first():
        return jsonify({"erro": "Esta pessoa já foi convidada para colaborar."}), 409

    conta = Usuario.query.filter(db.func.lower(Usuario.email) == email).first()
    colab = ColaboradorEvento(
        evento_id=evento.id,
        usuario_id=conta.id if conta else None,
        nome=nome,
        email=email,
        status="pendente",
        pode_editar=bool(dados.get("pode_editar", True)),
    )
    db.session.add(colab)
    db.session.flush()
    if conta:
        db.session.add(Notificacao(
            usuario_id=conta.id,
            tipo="convite-colaborador",
            titulo="Convite para colaborar",
            mensagem=f"{usuario.nome} convidou você para colaborar em {evento.nome_evento}.",
            evento_id=evento.id,
            referencia_id=colab.id,
        ))
    _log(evento.id, "colaborador_convidado", email)
    db.session.commit()
    resposta = colab.to_dict()
    resposta["link_convite"] = f"/convite.html?token={colab.token}"
    resposta["possui_conta"] = bool(conta)
    return jsonify(resposta), 201


@evento_detalhes_controller.post("/colaboracoes/<token>/responder")
@login_required
def responder_colaboracao(token):
    usuario = usuario_atual()
    colab = ColaboradorEvento.query.filter_by(token=token).first()
    if not colab:
        return jsonify({"erro": "Convite de colaboração inválido."}), 404
    if colab.email.lower() != usuario.email.lower():
        return jsonify({"erro": "Este convite foi enviado para outro e-mail."}), 403
    dados = request.get_json(silent=True) or {}
    status = dados.get("status")
    if status not in {"ativo", "recusado"}:
        return jsonify({"erro": "Resposta inválida."}), 400
    colab.usuario_id = usuario.id
    colab.nome = usuario.nome
    colab.status = status
    db.session.add(Notificacao(
        usuario_id=colab.evento.owner_id,
        tipo="colaborador-resposta",
        titulo="Resposta de colaborador",
        mensagem=f"{usuario.nome} {'aceitou' if status == 'ativo' else 'recusou'} colaborar em {colab.evento.nome_evento}.",
        evento_id=colab.evento_id,
        referencia_id=colab.id,
    ))
    _log(colab.evento_id, "colaborador_respondeu", f"{usuario.email}: {status}")
    db.session.commit()
    return jsonify(colab.to_dict())


@evento_detalhes_controller.put("/eventos/<int:evento_id>/colaboradores/<int:colaborador_id>")
@login_required
def atualizar_colaborador(evento_id, colaborador_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    colab = ColaboradorEvento.query.filter_by(id=colaborador_id, evento_id=evento.id).first()
    if not colab:
        return jsonify({"erro": "Colaborador não encontrado."}), 404
    dados = request.get_json(silent=True) or {}

    # O dono pode alterar permissões/status; o próprio colaborador pode responder.
    if evento.owner_id == usuario.id:
        if "pode_editar" in dados:
            colab.pode_editar = bool(dados["pode_editar"])
        if dados.get("status") in {"pendente", "ativo", "recusado"}:
            colab.status = dados["status"]
    elif colab.usuario_id == usuario.id and dados.get("status") in {"ativo", "recusado"}:
        colab.status = dados["status"]
    else:
        return jsonify({"erro": "Você não pode alterar este colaborador."}), 403
    db.session.commit()
    return jsonify(colab.to_dict())


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/colaboradores/<int:colaborador_id>")
@login_required
def remover_colaborador(evento_id, colaborador_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    if evento.owner_id != usuario.id:
        return jsonify({"erro": "Somente a pessoa dona do evento pode remover colaboradores."}), 403
    colab = ColaboradorEvento.query.filter_by(id=colaborador_id, evento_id=evento.id).first()
    if not colab:
        return jsonify({"erro": "Colaborador não encontrado."}), 404
    db.session.delete(colab)
    _log(evento.id, "colaborador_removido", colab.email)
    db.session.commit()
    return "", 204


# ---------------- Lista compartilhada ----------------

@evento_detalhes_controller.post("/eventos/<int:evento_id>/itens-compartilhados")
@login_required
def criar_item_compartilhado(evento_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    nome = (dados.get("nome") or "").strip()
    if not nome:
        return jsonify({"erro": "Informe o item."}), 400
    existe = ItemCompartilhado.query.filter(
        ItemCompartilhado.evento_id == evento.id,
        db.func.lower(ItemCompartilhado.nome) == nome.lower(),
    ).first()
    if existe:
        return jsonify({"erro": "Este item já está na lista.", "item": existe.to_dict()}), 409
    item = ItemCompartilhado(evento_id=evento.id, nome=nome, quantidade=max(1, int(dados.get("quantidade") or 1)))
    db.session.add(item)
    _log(evento.id, "item_compartilhado_adicionado", nome)
    db.session.commit()
    return jsonify(item.to_dict()), 201


@evento_detalhes_controller.put("/eventos/<int:evento_id>/itens-compartilhados/<int:item_id>")
@login_required
def atualizar_item_compartilhado(evento_id, item_id):
    usuario, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    item = ItemCompartilhado.query.filter_by(id=item_id, evento_id=evento.id).first()
    if not item:
        return jsonify({"erro": "Item não encontrado."}), 404
    dados = request.get_json(silent=True) or {}
    if dados.get("liberar"):
        item.responsavel_tipo = None
        item.responsavel_id = None
        item.responsavel_nome = ""
    elif "responsavelNome" in dados or "responsavel_nome" in dados:
        item.responsavel_tipo = dados.get("responsavelTipo") or "usuario"
        item.responsavel_id = int(dados.get("responsavelId")) if str(dados.get("responsavelId") or "").isdigit() else None
        item.responsavel_nome = (dados.get("responsavelNome") or dados.get("responsavel_nome") or usuario.nome).strip()
    if "quantidade" in dados:
        item.quantidade = max(1, int(dados.get("quantidade") or 1))
    db.session.commit()
    return jsonify(item.to_dict())


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/itens-compartilhados/<int:item_id>")
@login_required
def remover_item_compartilhado(evento_id, item_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    item = ItemCompartilhado.query.filter_by(id=item_id, evento_id=evento.id).first()
    if not item:
        return jsonify({"erro": "Item não encontrado."}), 404
    db.session.delete(item)
    db.session.commit()
    return "", 204


# ---------------- Itens do catálogo ----------------

@evento_detalhes_controller.post("/eventos/<int:evento_id>/itens")
@login_required
def adicionar_item_evento(evento_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    chave = str(dados.get("produtoId") or dados.get("produto_id") or dados.get("nome") or "").strip()
    existente = ItemEvento.query.filter_by(evento_id=evento.id, chave_externa=chave).first() if chave else None
    if existente:
        existente.quantidade += max(1, int(dados.get("qtd") or 1))
        db.session.commit()
        return jsonify(existente.to_dict())
    produto_id = int(dados["produto_id"]) if str(dados.get("produto_id") or "").isdigit() else None
    produto = db.session.get(ProdutoFornecedor, produto_id) if produto_id else None
    item = ItemEvento(
        evento_id=evento.id,
        produto_id=produto.id if produto else None,
        chave_externa=chave or None,
        nome=(dados.get("nome") or (produto.nome if produto else "Item")).strip(),
        categoria=(dados.get("categoria") or (produto.categoria if produto else "outros")).strip(),
        preco=float(dados.get("preco") if dados.get("preco") is not None else (produto.preco if produto else 0)),
        quantidade=max(1, int(dados.get("qtd") or 1)),
    )
    db.session.add(item)
    db.session.flush()
    _log(evento.id, "item_catalogo_adicionado", item.nome)
    if float(evento.orcamento_evento or 0) > 0 and evento.total_despesas > float(evento.orcamento_evento or 0):
        db.session.add(Notificacao(
            usuario_id=evento.owner_id, tipo="orcamento-ultrapassado", titulo="Teto de gastos ultrapassado",
            mensagem=f"{evento.nome_evento} ultrapassou o orçamento após adicionar {item.nome}.", evento_id=evento.id, referencia_id=item.id,
        ))
    db.session.commit()
    return jsonify(item.to_dict()), 201


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/itens/<int:item_id>")
@login_required
def remover_item_evento(evento_id, item_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    item = ItemEvento.query.filter_by(id=item_id, evento_id=evento.id).first()
    if not item:
        return jsonify({"erro": "Item não encontrado."}), 404
    db.session.delete(item)
    db.session.commit()
    return "", 204


# ---------------- Tarefas / cronograma ----------------

@evento_detalhes_controller.post("/eventos/<int:evento_id>/tarefas")
@login_required
def adicionar_tarefa(evento_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    titulo = (dados.get("titulo") or "").strip()
    if not titulo:
        return jsonify({"erro": "Título da tarefa é obrigatório."}), 400
    prazo = None
    if dados.get("prazo"):
        prazo = datetime.fromisoformat(dados["prazo"])
    tarefa = TarefaEvento(evento_id=evento.id, titulo=titulo, prazo=prazo, ordem=len(evento.tarefas_rel))
    db.session.add(tarefa)
    _log(evento.id, "tarefa_adicionada", titulo)
    db.session.commit()
    return jsonify(tarefa.to_dict()), 201


@evento_detalhes_controller.put("/eventos/<int:evento_id>/tarefas/<int:tarefa_id>")
@login_required
def atualizar_tarefa(evento_id, tarefa_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    tarefa = TarefaEvento.query.filter_by(id=tarefa_id, evento_id=evento.id).first()
    if not tarefa:
        return jsonify({"erro": "Tarefa não encontrada."}), 404
    dados = request.get_json(silent=True) or {}
    if "feita" in dados:
        tarefa.feita = bool(dados["feita"])
    if "titulo" in dados:
        tarefa.titulo = (dados.get("titulo") or tarefa.titulo).strip()
    if "prazo" in dados:
        tarefa.prazo = datetime.fromisoformat(dados["prazo"]) if dados.get("prazo") else None
    _recalcular_progresso(evento)
    db.session.commit()
    return jsonify({"tarefa": tarefa.to_dict(), "progresso": evento.progresso})


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/tarefas/<int:tarefa_id>")
@login_required
def remover_tarefa(evento_id, tarefa_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    tarefa = TarefaEvento.query.filter_by(id=tarefa_id, evento_id=evento.id).first()
    if not tarefa:
        return jsonify({"erro": "Tarefa não encontrada."}), 404
    db.session.delete(tarefa)
    db.session.flush()
    _recalcular_progresso(evento)
    db.session.commit()
    return "", 204


# ---------------- Despesas / orçamento / rateio ----------------

@evento_detalhes_controller.get("/eventos/<int:evento_id>/despesas")
@login_required
def listar_despesas(evento_id):
    _, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    return jsonify([d.to_dict() for d in evento.despesas_rel])


@evento_detalhes_controller.post("/eventos/<int:evento_id>/despesas")
@login_required
def adicionar_despesa(evento_id):
    usuario, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    descricao = (dados.get("descricao") or "").strip()
    valor = float(dados.get("valor") or 0)
    if not descricao or valor <= 0:
        return jsonify({"erro": "Informe descrição e valor maior que zero."}), 400
    despesa = DespesaEvento(
        evento_id=evento.id,
        descricao=descricao,
        categoria=(dados.get("categoria") or "geral").strip(),
        valor=valor,
        comprovante=(dados.get("comprovante") or "").strip(),
        pago_por=(dados.get("pago_por") or usuario.nome).strip(),
    )
    db.session.add(despesa)
    db.session.flush()
    _log(evento.id, "despesa_adicionada", f"{descricao}: {valor:.2f}")
    if float(evento.orcamento_evento or 0) > 0 and evento.total_despesas > float(evento.orcamento_evento or 0):
        db.session.add(Notificacao(
            usuario_id=evento.owner_id, tipo="orcamento-ultrapassado", titulo="Teto de gastos ultrapassado",
            mensagem=f"{evento.nome_evento} ultrapassou o orçamento definido.", evento_id=evento.id, referencia_id=despesa.id,
        ))
    db.session.commit()
    return jsonify({"despesa": despesa.to_dict(), "resumo": _resumo_financeiro(evento)}), 201


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/despesas/<int:despesa_id>")
@login_required
def remover_despesa(evento_id, despesa_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    despesa = DespesaEvento.query.filter_by(id=despesa_id, evento_id=evento.id).first()
    if not despesa:
        return jsonify({"erro": "Despesa não encontrada."}), 404
    db.session.delete(despesa)
    db.session.commit()
    return "", 204



@evento_detalhes_controller.post("/eventos/<int:evento_id>/despesas/<int:despesa_id>/comprovante")
@login_required
def enviar_comprovante_despesa(evento_id, despesa_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    despesa = DespesaEvento.query.filter_by(id=despesa_id, evento_id=evento.id).first()
    if not despesa:
        return jsonify({"erro": "Despesa não encontrada."}), 404
    arquivo = request.files.get("comprovante")
    if not arquivo or not arquivo.filename:
        return jsonify({"erro": "Selecione um comprovante."}), 400
    extensao = arquivo.filename.rsplit(".", 1)[-1].lower() if "." in arquivo.filename else ""
    if extensao not in {"png", "jpg", "jpeg", "webp", "gif", "pdf"}:
        return jsonify({"erro": "Comprovante deve ser imagem ou PDF."}), 400
    uploads = Path(os.getenv("UPLOAD_FOLDER", Path(__file__).resolve().parents[1] / "uploads"))
    uploads.mkdir(parents=True, exist_ok=True)
    nome_final = f"comprovante_evento{evento.id}_{despesa.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}_{secure_filename(arquivo.filename)}"
    arquivo.save(uploads / nome_final)
    despesa.comprovante = nome_final
    _log(evento.id, "comprovante_enviado", despesa.descricao)
    db.session.commit()
    return jsonify(despesa.to_dict())

def _resumo_financeiro(evento):
    total = evento.total_despesas
    orcamento = float(evento.orcamento_evento or 0)
    restante = orcamento - total
    ativos = [c for c in evento.colaboradores_rel if c.status == "ativo"]
    qtd = 1 + len(ativos)
    if evento.rateio_modo == "fixo":
        por_pessoa = float(evento.rateio_valor_fixo or 0)
    else:
        por_pessoa = total / qtd if qtd else 0
    return {
        "orcamento": orcamento,
        "total": total,
        "restante": restante,
        "percentual": round((total / orcamento) * 100, 1) if orcamento > 0 else 0,
        "estourou_orcamento": bool(orcamento > 0 and total > orcamento),
        "rateio": {
            "ativo": evento.rateio_ativo,
            "modo": evento.rateio_modo,
            "valorFixo": float(evento.rateio_valor_fixo or 0),
            "qtdPessoas": qtd,
            "valorPorPessoa": por_pessoa,
            "pessoas": ([{"id": "organizador", "nome": evento.owner.nome, "valor": por_pessoa, "organizador": True}] + [
                {"id": str(c.id), "nome": c.nome, "valor": por_pessoa, "organizador": False} for c in ativos
            ]),
        },
    }


@evento_detalhes_controller.get("/eventos/<int:evento_id>/resumo")
@login_required
def resumo_evento(evento_id):
    _, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro

    confirmados = evento.confirmados
    base = max(confirmados, evento.convidados_estimados or 0)
    # Estimativas simples e transparentes; são sugestões, não compras automáticas.
    suprimentos = {
        "pessoas_consideradas": base,
        "comida_kg": round(base * 0.45, 1),
        "bebidas_litros": round(base * 0.8, 1),
        "gelo_kg": round(base * 0.2, 1),
        "doces_unidades": math.ceil(base * 5),
    }
    restricoes = [c.restricoes_alimentares for c in evento.convidados_rel if c.restricoes_alimentares]

    perfil_custos = {
        "casamento": {"alimentacao_por_pessoa": 140, "bebida_por_pessoa": 45, "decoracao": 8000, "local": 6000},
        "aniversario": {"alimentacao_por_pessoa": 80, "bebida_por_pessoa": 25, "decoracao": 2500, "local": 2500},
        "corporativo": {"alimentacao_por_pessoa": 110, "bebida_por_pessoa": 30, "decoracao": 3000, "local": 4500},
    }
    custo = perfil_custos.get((evento.tipo_evento or "").lower(), {"alimentacao_por_pessoa": 90, "bebida_por_pessoa": 30, "decoracao": 3000, "local": 3000})
    orcamento_base = {
        "alimentacao": round(base * custo["alimentacao_por_pessoa"], 2),
        "bebidas": round(base * custo["bebida_por_pessoa"], 2),
        "decoracao": float(custo["decoracao"]),
        "local": float(custo["local"]),
    }
    orcamento_base["total_estimado"] = round(sum(orcamento_base.values()), 2)

    consulta_fornecedores = Fornecedor.query.filter_by(ativo=True)
    fornecedores = consulta_fornecedores.order_by(Fornecedor.avaliacao.desc(), Fornecedor.nome_empresa.asc()).limit(6).all()
    consulta_locais = Fornecedor.query.filter(Fornecedor.ativo.is_(True), Fornecedor.tipo_servico.ilike("%espa%"))
    if evento.cidade_evento:
        consulta_locais = consulta_locais.filter(Fornecedor.cidade.ilike(f"%{evento.cidade_evento}%"))
    locais = consulta_locais.order_by(Fornecedor.avaliacao.desc()).limit(6).all()
    return jsonify({
        "financeiro": _resumo_financeiro(evento),
        "convidados": {
            "estimados": evento.convidados_estimados,
            "confirmados": confirmados,
            "pendentes": sum(1 for c in evento.convidados_rel if c.status == "pendente"),
            "recusados": sum(1 for c in evento.convidados_rel if c.status == "recusado"),
            "restricoes_alimentares": restricoes,
        },
        "suprimentos": suprimentos,
        "orcamento_base": orcamento_base,
        "mapa_url": "https://www.google.com/maps/search/?api=1&query=" + (evento.endereco_evento or evento.cidade_evento or "espaço para eventos").replace(" ", "+"),
        "fornecedores_recomendados": [f.to_dict(incluir_produtos=False) for f in fornecedores],
        "locais_recomendados": [f.to_dict(incluir_produtos=True) for f in locais],
    })


@evento_detalhes_controller.put("/eventos/<int:evento_id>/rateio")
@login_required
def definir_rateio(evento_id):
    _, evento = acesso_evento(evento_id, exigir_edicao=True)
    erro = _erro(evento)
    if erro:
        return erro
    dados = request.get_json(silent=True) or {}
    if "ativo" in dados:
        evento.rateio_ativo = bool(dados["ativo"])
    if dados.get("modo") in {"auto", "fixo"}:
        evento.rateio_modo = dados["modo"]
    if "valorFixo" in dados:
        evento.rateio_valor_fixo = max(0, float(dados.get("valorFixo") or 0))
    db.session.commit()
    return jsonify(_resumo_financeiro(evento)["rateio"])


# ---------------- Logs ----------------

@evento_detalhes_controller.get("/eventos/<int:evento_id>/logs")
@login_required
def listar_logs(evento_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    if evento.owner_id != usuario.id:
        return jsonify({"erro": "Histórico disponível apenas para a pessoa dona do evento."}), 403
    logs = LogEvento.query.filter_by(evento_id=evento.id).order_by(LogEvento.criado_em.desc()).limit(100).all()
    return jsonify([l.to_dict() for l in logs])


# ---------------- Galeria ----------------

ALLOWED_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif"}
ALLOWED_RECEIPT_EXTENSIONS = ALLOWED_IMAGE_EXTENSIONS | {"pdf"}


@evento_detalhes_controller.post("/eventos/<int:evento_id>/fotos")
@login_required
def enviar_foto(evento_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    arquivo = request.files.get("foto")
    if not arquivo or not arquivo.filename:
        return jsonify({"erro": "Selecione uma imagem."}), 400
    extensao = arquivo.filename.rsplit(".", 1)[-1].lower() if "." in arquivo.filename else ""
    if extensao not in ALLOWED_IMAGE_EXTENSIONS:
        return jsonify({"erro": "Formato de imagem não suportado."}), 400

    uploads = Path(os.getenv("UPLOAD_FOLDER", Path(__file__).resolve().parents[1] / "uploads"))
    uploads.mkdir(parents=True, exist_ok=True)
    nome_seguro = secure_filename(arquivo.filename)
    nome_final = f"evento{evento.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}_{nome_seguro}"
    arquivo.save(uploads / nome_final)
    foto = FotoEvento(
        evento_id=evento.id,
        usuario_id=usuario.id,
        arquivo=nome_final,
        legenda=(request.form.get("legenda") or "").strip(),
    )
    db.session.add(foto)
    db.session.commit()
    return jsonify(foto.to_dict()), 201


@evento_detalhes_controller.delete("/eventos/<int:evento_id>/fotos/<int:foto_id>")
@login_required
def remover_foto(evento_id, foto_id):
    usuario, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    foto = FotoEvento.query.filter_by(id=foto_id, evento_id=evento.id).first()
    if not foto:
        return jsonify({"erro": "Foto não encontrada."}), 404
    if evento.owner_id != usuario.id and foto.usuario_id != usuario.id:
        return jsonify({"erro": "Você não pode remover esta foto."}), 403
    db.session.delete(foto)
    db.session.commit()
    return "", 204


# ---------------- Relatórios ----------------

def _linhas_relatorio(evento):
    return [
        ["Evento", evento.nome_evento],
        ["Tipo", evento.tipo_evento],
        ["Data", evento.data_evento.strftime("%d/%m/%Y")],
        ["Local", evento.endereco_evento or "-"],
        ["Orçamento", f"R$ {float(evento.orcamento_evento or 0):.2f}"],
        ["Gasto atual", f"R$ {evento.total_despesas:.2f}"],
        ["Confirmados", str(evento.confirmados)],
    ]


@evento_detalhes_controller.get("/eventos/<int:evento_id>/relatorio.pdf")
@login_required
def relatorio_pdf(evento_id):
    _, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=f"Relatório - {evento.nome_evento}")
    estilos = getSampleStyleSheet()
    elementos = [Paragraph("Momentus - Relatório do Evento", estilos["Title"]), Spacer(1, 12)]
    tabela = Table(_linhas_relatorio(evento), colWidths=[120, 340])
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#2b1a4d")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#ddd6ee")),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    elementos += [tabela, Spacer(1, 16), Paragraph("Convidados", estilos["Heading2"])]
    if evento.convidados_rel:
        dados = [["Nome", "Status", "Restrições"]] + [[c.nome, c.status, c.restricoes_alimentares or "-"] for c in evento.convidados_rel]
        elementos.append(Table(dados, repeatRows=1, colWidths=[180, 90, 190]))
    else:
        elementos.append(Paragraph("Nenhum convidado cadastrado.", estilos["BodyText"]))
    elementos += [Spacer(1, 16), Paragraph("Despesas", estilos["Heading2"])]
    if evento.despesas_rel:
        dados = [["Descrição", "Categoria", "Valor"]] + [[d.descricao, d.categoria, f"R$ {float(d.valor):.2f}"] for d in evento.despesas_rel]
        elementos.append(Table(dados, repeatRows=1, colWidths=[230, 130, 110]))
    else:
        elementos.append(Paragraph("Nenhuma despesa cadastrada.", estilos["BodyText"]))
    elementos += [Spacer(1, 16), Paragraph("Itens do evento", estilos["Heading2"])]
    if evento.itens_evento_rel:
        dados = [["Item", "Categoria", "Qtd.", "Valor"]] + [[i.nome, i.categoria or "-", str(i.quantidade or 1), f"R$ {float(i.preco or 0) * (i.quantidade or 1):.2f}"] for i in evento.itens_evento_rel]
        elementos.append(Table(dados, repeatRows=1, colWidths=[190, 120, 60, 100]))
    else:
        elementos.append(Paragraph("Nenhum item de catálogo adicionado.", estilos["BodyText"]))
    elementos += [Spacer(1, 16), Paragraph("Lista compartilhada", estilos["Heading2"])]
    if evento.itens_compartilhados_rel:
        dados = [["Item", "Qtd.", "Responsável"]] + [[i.nome, str(i.quantidade or 1), i.responsavel_nome or "-"] for i in evento.itens_compartilhados_rel]
        elementos.append(Table(dados, repeatRows=1, colWidths=[230, 70, 170]))
    else:
        elementos.append(Paragraph("Nenhum item compartilhado cadastrado.", estilos["BodyText"]))
    doc.build(elementos)
    buffer.seek(0)
    return send_file(buffer, mimetype="application/pdf", as_attachment=True, download_name=f"momentus-evento-{evento.id}.pdf")


@evento_detalhes_controller.get("/eventos/<int:evento_id>/relatorio.xlsx")
@login_required
def relatorio_xlsx(evento_id):
    _, evento = acesso_evento(evento_id)
    erro = _erro(evento)
    if erro:
        return erro
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    for linha in _linhas_relatorio(evento):
        ws.append(linha)
    ws2 = wb.create_sheet("Convidados")
    ws2.append(["Nome", "E-mail", "Telefone", "Status", "Restrições"])
    for c in evento.convidados_rel:
        ws2.append([c.nome, c.email, c.telefone, c.status, c.restricoes_alimentares])
    ws3 = wb.create_sheet("Despesas")
    ws3.append(["Descrição", "Categoria", "Valor", "Pago por", "Comprovante"])
    for d in evento.despesas_rel:
        ws3.append([d.descricao, d.categoria, float(d.valor), d.pago_por, d.comprovante])
    ws4 = wb.create_sheet("Itens")
    ws4.append(["Item", "Categoria", "Quantidade", "Preço unitário", "Subtotal"])
    for i in evento.itens_evento_rel:
        ws4.append([i.nome, i.categoria, i.quantidade, float(i.preco or 0), float(i.preco or 0) * (i.quantidade or 1)])
    ws5 = wb.create_sheet("Lista compartilhada")
    ws5.append(["Item", "Quantidade", "Responsável"])
    for i in evento.itens_compartilhados_rel:
        ws5.append([i.nome, i.quantidade, i.responsavel_nome])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"momentus-evento-{evento.id}.xlsx")
