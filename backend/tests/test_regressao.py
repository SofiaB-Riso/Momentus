"""Regressões de API, transação e contratos introduzidas pela refatoração."""

import ast
import io
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["DATABASE_URL"] = "sqlite://"
from app import create_app
from application.erros import DadosInvalidos
from models import (
    ColaboradorEvento,
    Evento,
    ItemCompartilhado,
    LogEvento,
    Usuario,
    db,
)
from services.caso_uso import CasoUso
from services.financeiro.resumo_financeiro_service import resumo_financeiro


def imagem():
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), "purple").save(buffer, "PNG")
    buffer.seek(0)
    return buffer


class RegressaoTest(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.TemporaryDirectory(prefix="momentus-regressao-")
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "chave-teste",
                "SQLALCHEMY_DATABASE_URI": "sqlite:///"
                + str(Path(self.pasta.name) / "teste.db"),
                "UPLOAD_FOLDER": str(Path(self.pasta.name) / "arquivos"),
                "LIVE_INTERVAL": 0,
                "LIVE_ITERATIONS": 1,
            }
        )
        self.cliente = self.app.test_client()
        r = self.cliente.post(
            "/auth/cadastro",
            json={
                "nome": "Organizador",
                "email": "dono@teste.example",
                "senha": "senha-teste",
            },
        )
        self.assertEqual(r.status_code, 201)
        self.usuario_id = r.get_json()["usuario"]["id"]
        self.payload = {
            "nome_evento": "Evento de regressão",
            "tipo_evento": "aniversario",
            "data_evento": (date.today() + timedelta(days=1)).isoformat(),
            "cidade_evento": "São Paulo",
            "orcamento_evento": 1000,
        }
        r = self.cliente.post("/eventos", json=self.payload)
        self.assertEqual(r.status_code, 201, r.get_json())
        self.evento_id = r.get_json()["id"]
        self.base = f"/eventos/{self.evento_id}"

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        self.pasta.cleanup()

    def test_login_logout_perfil_e_email_duplicado(self):
        self.assertEqual(self.cliente.post("/auth/logout").status_code, 204)
        self.assertEqual(self.cliente.get("/auth/me").status_code, 401)
        self.assertEqual(
            self.cliente.post(
                "/auth/login",
                json={"email": "dono@teste.example", "senha": "incorreta"},
            ).status_code,
            401,
        )
        self.assertEqual(
            self.cliente.post(
                "/auth/login",
                json={"email": " DONO@TESTE.EXAMPLE ", "senha": "senha-teste"},
            ).status_code,
            200,
        )
        r = self.cliente.put(
            "/auth/perfil", json={"nome": "Nome atualizado", "telefone": "31999999999"}
        )
        self.assertEqual(r.get_json()["usuario"]["nome"], "Nome atualizado")
        outro = self.app.test_client()
        self.assertEqual(
            outro.post(
                "/auth/cadastro",
                json={
                    "nome": "Outro",
                    "email": "dono@teste.example",
                    "senha": "senha-teste",
                },
            ).status_code,
            409,
        )

    def test_crud_evento_e_consulta_por_tipo(self):
        r = self.cliente.get("/eventos/por-tipo?tipo_evento=aniversario")
        self.assertEqual([e["id"] for e in r.get_json()], [self.evento_id])
        self.assertEqual(self.cliente.get("/eventos/por-tipo").status_code, 400)
        self.assertEqual(
            self.cliente.put(
                self.base, json={"nome_evento": "Evento alterado"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.cliente.get(self.base).get_json()["nome_evento"], "Evento alterado"
        )
        self.assertEqual(self.cliente.delete(self.base).status_code, 204)
        self.assertEqual(self.cliente.get(self.base).status_code, 404)
        with self.app.app_context():
            self.assertIsNone(Evento.buscar_por_id(self.evento_id))
            self.assertEqual(
                LogEvento.query.filter_by(evento_id=self.evento_id).count(), 0
            )

    def test_erro_no_historico_reverte_criacao_e_tarefas(self):
        with patch.object(CasoUso, "log", side_effect=DadosInvalidos("Falha simulada")):
            r = self.cliente.post(
                "/eventos", json={**self.payload, "nome_evento": "Não deve persistir"}
            )
            self.assertEqual(r.status_code, 400)
        with self.app.app_context():
            self.assertEqual(Evento.query.count(), 1)
            self.assertIsNone(Evento.buscar_um(nome_evento="Não deve persistir"))

    def test_payloads_invalidos_nao_geram_erro_500(self):
        for payload in (
            ["lista"],
            {**self.payload, "nome_evento": 42},
            {**self.payload, "data_evento": "data inválida"},
            {**self.payload, "privacidade_convidados": "invalida"},
        ):
            r = self.cliente.post("/eventos", json=payload)
            self.assertEqual(r.status_code, 400, (payload, r.get_json()))
        for payload in (
            {"titulo": "Tarefa", "prazo": "data inválida"},
            {"titulo": "Tarefa", "feita": "false"},
        ):
            self.assertEqual(
                self.cliente.post(self.base + "/tarefas", json=payload).status_code, 400
            )

    def test_horario_valido_criacao_edicao_e_opcional(self):
        for horario, esperado in [
            ("00:00", "00:00"),
            ("23:59", "23:59"),
            (" 09:05 ", "09:05"),
            ("", ""),
            (None, ""),
        ]:
            with self.subTest(horario=horario):
                r = self.cliente.post(
                    "/eventos", json={**self.payload, "hora_evento": horario}
                )
                self.assertEqual(r.status_code, 201, r.get_json())
                self.assertEqual(r.get_json()["hora_evento"], esperado)
                r = self.cliente.put(self.base, json={"hora_evento": horario})
                self.assertEqual(r.status_code, 200, r.get_json())
                self.assertEqual(r.get_json()["hora_evento"], esperado)
        self.assertEqual(
            self.cliente.put(self.base, json={"hora_evento": "14:37"}).status_code, 200
        )
        self.assertEqual(
            self.cliente.put(
                self.base, json={"nome_evento": "Sem alterar hora"}
            ).status_code,
            200,
        )
        self.assertEqual(self.cliente.get(self.base).get_json()["hora_evento"], "14:37")

    def test_horario_invalido_rejeitado_sem_persistir(self):
        self.cliente.put(self.base, json={"hora_evento": "15:30"})
        antes = self.cliente.get(self.base).get_json()
        quantidade = len(self.cliente.get("/eventos").get_json())
        for horario in [
            "24:00",
            "12:60",
            "-1:00",
            "9:05",
            "1500",
            "15:00:00",
            "abc",
            "１２:００",
            1500,
            [],
            {},
        ]:
            with self.subTest(horario=horario):
                r = self.cliente.post(
                    "/eventos", json={**self.payload, "hora_evento": horario}
                )
                self.assertEqual(r.status_code, 400, r.get_json())
                r = self.cliente.put(
                    self.base,
                    json={"hora_evento": horario, "nome_evento": "Não persistir"},
                )
                self.assertEqual(r.status_code, 400, r.get_json())
                depois = self.cliente.get(self.base).get_json()
                self.assertEqual(depois["hora_evento"], antes["hora_evento"])
                self.assertEqual(depois["nome_evento"], antes["nome_evento"])
        self.assertEqual(len(self.cliente.get("/eventos").get_json()), quantidade)

    def test_upload_galeria_pasta_configurada_e_remocao(self):
        r = self.cliente.post(
            self.base + "/fotos",
            data={"foto": (imagem(), "foto.png"), "legenda": "Foto compartilhada"},
        )
        self.assertEqual(r.status_code, 201, r.get_json())
        foto = r.get_json()
        caminho = Path(self.app.config["UPLOAD_FOLDER"]) / foto["arquivo"]
        self.assertTrue(caminho.exists())
        resposta = self.cliente.get(foto["url"])
        self.assertEqual(resposta.status_code, 200)
        self.assertTrue(resposta.data.startswith(b"\x89PNG"))
        resposta.close()
        self.assertEqual(
            self.cliente.delete(self.base + "/fotos/" + foto["id"]).status_code, 204
        )
        self.assertFalse(caminho.exists())
        self.assertEqual(
            self.cliente.post(
                self.base + "/fotos",
                data={"foto": (io.BytesIO(b"texto falso"), "foto.png")},
            ).status_code,
            400,
        )

    def test_upload_falha_no_banco_remove_arquivo_parcial(self):
        with patch.object(CasoUso, "log", side_effect=DadosInvalidos("Falha simulada")):
            r = self.cliente.post(
                self.base + "/fotos", data={"foto": (imagem(), "foto.png")}
            )
            self.assertEqual(r.status_code, 400)
        self.assertEqual(list(Path(self.app.config["UPLOAD_FOLDER"]).iterdir()), [])
        self.assertEqual(self.cliente.get(self.base).get_json()["fotos"], [])

    def test_comprovante_substitui_arquivo_e_valida_conteudo(self):
        d = self.cliente.post(
            self.base + "/despesas", json={"descricao": "Bolo", "valor": "12.35"}
        ).get_json()["despesa"]
        url = self.base + "/despesas/" + d["id"] + "/comprovante"
        primeiro = self.cliente.post(
            url, data={"comprovante": (imagem(), "recibo.png")}
        ).get_json()
        antigo = Path(self.app.config["UPLOAD_FOLDER"]) / primeiro["comprovante"]
        self.assertTrue(antigo.exists())
        segundo = self.cliente.post(url, data={"comprovante": (imagem(), "novo.png")})
        self.assertEqual(segundo.status_code, 200)
        self.assertFalse(antigo.exists())
        self.assertEqual(
            self.cliente.post(
                url, data={"comprovante": (io.BytesIO(b"nao-pdf"), "recibo.pdf")}
            ).status_code,
            400,
        )

    def test_relatorios_pdf_excel_incluem_dados_e_texto_seguro(self):
        self.cliente.post(
            self.base + "/convidados",
            json={"nome": "=2+2", "restricoes": "Alergia <castanhas> & leite"},
        )
        self.cliente.post(
            self.base + "/despesas", json={"descricao": "Bolo", "valor": "12.35"}
        )
        pdf = self.cliente.get(self.base + "/relatorio.pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b"%PDF-"))
        planilha = self.cliente.get(self.base + "/relatorio.xlsx")
        self.assertEqual(planilha.status_code, 200)
        wb = load_workbook(io.BytesIO(planilha.data))
        self.assertEqual(
            set(wb.sheetnames),
            {"Resumo", "Convidados", "Despesas", "Itens", "Lista compartilhada"},
        )
        self.assertEqual(wb["Convidados"]["A2"].value, "'=2+2")
        self.assertEqual(wb["Despesas"]["C2"].value, 12.35)
        wb.close()
        pdf.close()
        planilha.close()

    def test_contratos_aceitam_provedor_e_exportador_substitutos(self):
        class ProvedorAlternativo:
            def buscar(self, categoria, **filtros):
                return [
                    {
                        "nome_empresa": "Fornecedor alternativo",
                        "avaliacao": 4.9,
                        "total_avaliacoes": 3,
                    }
                ]

        class ExportadorTexto:
            def exportar(self, dados):
                return dados["nome"].encode()

        contexto = self.app.extensions["momentus"]
        contexto.provedor = ProvedorAlternativo()
        self.assertEqual(
            self.cliente.get(self.base + "/recomendacoes").get_json()["resultados"][0][
                "nome_empresa"
            ],
            "Fornecedor alternativo",
        )
        contexto.exportadores["txt"] = (ExportadorTexto(), "text/plain")
        with self.app.app_context():
            arquivo = contexto.relatorio("txt", self.usuario_id).executar(
                self.evento_id
            )
            self.assertEqual(arquivo.conteudo.decode(), self.payload["nome_evento"])

    def test_total_financeiro_consistente_apos_incluir_e_remover_itens(self):
        item = self.cliente.post(
            self.base + "/itens",
            json={"nome": "Item com centavos", "preco": "0.10", "qtd": 3},
        )
        self.assertEqual(item.status_code, 201)
        despesa = self.cliente.post(
            self.base + "/despesas",
            json={"descricao": "Despesa com centavos", "valor": "0.07"},
        )
        self.assertEqual(despesa.status_code, 201)
        evento = self.cliente.get(self.base).get_json()
        self.assertEqual(evento["total_despesas"], 0.37)
        self.assertEqual(evento["financeiro"]["total"], 0.37)
        self.assertEqual(evento["financeiro"]["rateio"]["totalDespesas"], 0.37)
        self.assertEqual(
            self.cliente.delete(
                self.base + "/itens/" + str(item.json["id"])
            ).status_code,
            204,
        )
        evento = self.cliente.get(self.base).get_json()
        self.assertEqual(evento["total_despesas"], 0.07)
        self.assertEqual(evento["financeiro"]["total"], 0.07)

    def test_rateio_aceita_estrategia_por_contrato(self):
        class EstrategiaAlternativa:
            def calcular(self, total, pessoas, valor_fixo):
                return [7] * pessoas

        with self.app.app_context():
            evento = Evento.buscar_por_id(self.evento_id)
            resultado = resumo_financeiro(evento, {"auto": EstrategiaAlternativa()})
            self.assertEqual(resultado["rateio"]["valorPorPessoa"], 7)

    def test_exclusao_conta_remove_eventos_e_libera_itens_de_colaboracao(self):
        outro = self.app.test_client()
        outro.post(
            "/auth/cadastro",
            json={
                "nome": "Outro",
                "email": "outro@teste.example",
                "senha": "senha-teste",
            },
        )
        evento = outro.post("/eventos", json=self.payload).get_json()
        base = f"/eventos/{evento['id']}"
        convite = outro.post(
            base + "/colaboradores",
            json={"nome": "Organizador", "email": "dono@teste.example"},
        ).get_json()
        self.cliente.post(
            "/colaboracoes/" + convite["token"] + "/responder", json={"status": "ativo"}
        )
        item = outro.post(
            base + "/itens-compartilhados", json={"nome": "Bebida"}
        ).get_json()
        self.cliente.put(
            base + "/itens-compartilhados/" + item["id"],
            json={"responsavelTipo": "usuario", "responsavelNome": "Organizador"},
        )
        self.assertEqual(self.cliente.delete("/auth/conta").status_code, 204)
        self.assertEqual(self.cliente.get("/auth/me").status_code, 401)
        with self.app.app_context():
            self.assertIsNone(Usuario.buscar_por_id(self.usuario_id))
            self.assertIsNone(Evento.buscar_por_id(self.evento_id))
            self.assertFalse(
                ItemCompartilhado.buscar_por_id(item["id"]).responsavel_nome
            )
            self.assertEqual(
                ColaboradorEvento.buscar_por_id(convite["id"]).status, "recusado"
            )

    def test_consulta_eventos_aceita_repositorio_substituto(self):
        class RepositorioAlternativo:
            def buscar_por_tipo_evento(self, tipo_evento, usuario_id):
                self.parametros = (tipo_evento, usuario_id)
                return []

        repositorio = RepositorioAlternativo()
        self.app.extensions["momentus"].busca_eventos = repositorio
        r = self.cliente.get("/eventos/por-tipo?tipo_evento=aniversario")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json(), [])
        self.assertEqual(repositorio.parametros, ("aniversario", self.usuario_id))

    def test_exclusao_despesa_evento_e_conta_remove_arquivos(self):
        despesa = self.cliente.post(
            self.base + "/despesas", json={"descricao": "Bolo", "valor": 10}
        ).get_json()["despesa"]
        url = self.base + "/despesas/" + despesa["id"]
        r = self.cliente.post(
            url + "/comprovante", data={"comprovante": (imagem(), "recibo.png")}
        )
        comprovante = (
            Path(self.app.config["UPLOAD_FOLDER"]) / r.get_json()["comprovante"]
        )
        self.assertTrue(comprovante.exists())
        self.assertEqual(self.cliente.delete(url).status_code, 204)
        self.assertFalse(comprovante.exists())
        r = self.cliente.post(
            self.base + "/fotos", data={"foto": (imagem(), "foto.png")}
        )
        foto = Path(self.app.config["UPLOAD_FOLDER"]) / r.get_json()["arquivo"]
        self.assertEqual(self.cliente.delete(self.base).status_code, 204)
        self.assertFalse(foto.exists())
        evento = self.cliente.post("/eventos", json=self.payload).get_json()
        r = self.cliente.post(
            f"/eventos/{evento['id']}/fotos", data={"foto": (imagem(), "foto.png")}
        )
        foto = Path(self.app.config["UPLOAD_FOLDER"]) / r.get_json()["arquivo"]
        self.assertEqual(self.cliente.delete("/auth/conta").status_code, 204)
        self.assertFalse(foto.exists())

    def test_remover_ultima_tarefa_zera_progresso(self):
        tarefas = self.cliente.get(self.base).get_json()["tarefas"]
        for tarefa in tarefas:
            self.assertEqual(
                self.cliente.put(
                    self.base + "/tarefas/" + tarefa["id"], json={"feita": True}
                ).status_code,
                200,
            )
        self.assertEqual(self.cliente.get(self.base).get_json()["progresso"], 100)
        for tarefa in tarefas:
            self.assertEqual(
                self.cliente.delete(self.base + "/tarefas/" + tarefa["id"]).status_code,
                204,
            )
        self.assertEqual(self.cliente.get(self.base).get_json()["progresso"], 0)

    @unittest.skipUnless(
        hasattr(__import__("time"), "tzset"), "Requer alteração de fuso do processo"
    )
    def test_lembretes_deduplicam_dia_local_com_timestamps_utc(self):
        import time
        from datetime import datetime, timezone

        from models import Notificacao
        from repositories.notificacao_repository import NotificacaoRepository

        for fuso in ("Australia/Brisbane", "America/Sao_Paulo"):
            try:
                with patch.dict(os.environ, {"TZ": fuso}):
                    time.tzset()
                    hoje = date.today()
                    inicio = (
                        datetime.combine(hoje, datetime.min.time())
                        .astimezone(timezone.utc)
                        .replace(tzinfo=None)
                    )
                    with self.app.app_context():
                        aviso = Notificacao(
                            usuario_id=self.usuario_id,
                            tipo=fuso,
                            titulo="Hoje",
                            mensagem="Hoje",
                            criado_em=datetime.now(timezone.utc).replace(tzinfo=None),
                        ).salvar()
                        self.assertEqual(
                            NotificacaoRepository.existente_hoje(
                                hoje, usuario_id=self.usuario_id, tipo=fuso
                            ).id,
                            aviso.id,
                        )
                        aviso.atualizar(criado_em=inicio - timedelta(seconds=1))
                        self.assertIsNone(
                            NotificacaoRepository.existente_hoje(
                                hoje, usuario_id=self.usuario_id, tipo=fuso
                            )
                        )
                        aviso.atualizar(criado_em=inicio + timedelta(days=1))
                        self.assertIsNone(
                            NotificacaoRepository.existente_hoje(
                                hoje, usuario_id=self.usuario_id, tipo=fuso
                            )
                        )
            finally:
                time.tzset()

    def test_rascunho_persiste_e_finaliza_sem_duplicar_evento_ou_tarefas(self):
        r = self.cliente.post("/eventos", json={**self.payload, "status": "rascunho"})
        self.assertEqual(r.status_code, 201)
        rascunho = r.get_json()
        base = f"/eventos/{rascunho['id']}"
        self.assertEqual(rascunho["status"], "rascunho")
        tarefas = [t["id"] for t in rascunho["tarefas"]]
        quantidade = len(self.cliente.get("/eventos").get_json())
        r = self.cliente.put(
            base, json={"nome_evento": "Rascunho revisado", "status": "rascunho"}
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["status"], "rascunho")
        r = self.cliente.put(base, json={"status": "confirmado"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["status"], "confirmado")
        self.assertEqual([t["id"] for t in r.get_json()["tarefas"]], tarefas)
        self.assertEqual(len(self.cliente.get("/eventos").get_json()), quantidade)

    def test_tipo_personalizado_persiste_e_pode_ser_editado(self):
        resposta = self.cliente.post(
            "/eventos", json={**self.payload, "tipo_evento": "Festa Junina"}
        )
        self.assertEqual(resposta.status_code, 201)
        evento_id = resposta.get_json()["id"]
        url = f"/eventos/{evento_id}"
        self.assertEqual(
            self.cliente.get(url).get_json()["tipo_evento"], "Festa Junina"
        )
        self.assertEqual(
            self.cliente.put(url, json={"tipo_evento": "Bodas de Prata"}).status_code,
            200,
        )
        filtrados = self.cliente.get(
            "/eventos/por-tipo", query_string={"tipo_evento": "Bodas de Prata"}
        ).get_json()
        self.assertEqual([evento["id"] for evento in filtrados], [evento_id])
        self.assertEqual(
            self.cliente.put(url, json={"tipo_evento": "a" * 101}).status_code, 400
        )
        self.assertEqual(
            self.cliente.get(url).get_json()["tipo_evento"], "Bodas de Prata"
        )

    def test_controller_e_service_preservam_fronteiras_da_arquitetura(self):
        raiz = Path(__file__).resolve().parents[1]
        for pasta in ("controllers", "services"):
            for caminho in (raiz / pasta).rglob("*.py"):
                tree = ast.parse(caminho.read_text())
                for node in ast.walk(tree):
                    if isinstance(node, ast.Attribute):
                        self.assertNotEqual(node.attr, "query", str(caminho))
                        if isinstance(node.value, ast.Name) and node.value.id == "db":
                            self.fail(
                                "Banco acessado fora de Model/Repository: "
                                + str(caminho)
                            )
                    if pasta == "services" and isinstance(node, ast.ImportFrom):
                        self.assertFalse(
                            (node.module or "").startswith(("flask", "infrastructure")),
                            str(caminho),
                        )
        for caminho in (raiz / "models").glob("*.py"):
            for node in ast.walk(ast.parse(caminho.read_text())):
                if isinstance(node, ast.ImportFrom):
                    self.assertFalse(
                        (node.module or "").startswith("services"), str(caminho)
                    )


if __name__ == "__main__":
    unittest.main()
