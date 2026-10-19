import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


class ExportadorExcel:
    def exportar(self, dados):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for titulo, linhas in dados.items():
            if titulo == "nome":
                continue
            folha = workbook.create_sheet(titulo)
            for linha in linhas:
                folha.append(
                    [
                        "'" + v
                        if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                        else v
                        for v in linha
                    ]
                )
            for celula in folha[1]:
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor="2B1A4D")
            folha.freeze_panes = "A2"
            folha.auto_filter.ref = folha.dimensions
            for coluna in folha.columns:
                folha.column_dimensions[coluna[0].column_letter].width = min(
                    55, max(16, max(len(str(c.value or "")) for c in coluna) + 2)
                )
                for celula in coluna:
                    celula.alignment = Alignment(vertical="top", wrap_text=True)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
